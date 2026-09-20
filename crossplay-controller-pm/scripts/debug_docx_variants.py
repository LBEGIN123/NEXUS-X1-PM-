from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "_qa" / "variants"
OUT.mkdir(parents=True, exist_ok=True)


def load_builder():
    spec = importlib.util.spec_from_file_location("prd_builder", ROOT / "scripts" / "build_prd_docx.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


def save_minimal(name: str):
    doc = Document()
    doc.add_paragraph("hello")
    doc.save(OUT / name)


def save_footer(builder):
    doc = Document()
    builder.add_page_number(doc.sections[0].footer.paragraphs[0])
    doc.add_paragraph("hello")
    doc.save(OUT / "footer.docx")


def save_table(builder):
    doc = Document()
    builder.add_table(doc, [["A", "B"], ["C", "D"]])
    doc.save(OUT / "table.docx")


def save_headings(builder):
    doc = Document()
    builder.configure_styles(doc)
    builder.add_heading(doc, "Heading one", 1)
    builder.add_paragraph(doc, "Body paragraph")
    builder.add_bullet(doc, "Bullet item")
    builder.add_bullet(doc, "Number item", ordered=True)
    doc.save(OUT / "headings.docx")


def save_full_variant(builder, name: str, remove_tables=False, remove_footer=False):
    markdown = (ROOT / "docs" / "PRD.md").read_text(encoding="utf-8")
    blocks = builder.parse_markdown(markdown)
    doc = Document()
    builder.configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    if not remove_footer:
        builder.add_page_number(section.footer.paragraphs[0])
    title = next((block[2] for block in blocks if block[0] == "heading" and block[1] == 1), "PRD")
    builder.add_cover(doc, title)
    for block in blocks:
        if block[0] in ("heading", "paragraph", "bullet", "number", "table"):
            if block[0] == "heading":
                if block[1] == 1:
                    continue
                builder.add_heading(doc, block[2], min(block[1] - 1, 3))
            elif block[0] == "table":
                if not remove_tables:
                    builder.add_table(doc, block[1])
            elif block[0] == "bullet":
                builder.add_bullet(doc, block[1])
            elif block[0] == "number":
                builder.add_numbered_paragraph(doc, block[1], block[2])
            else:
                builder.add_paragraph(doc, block[1])
    doc.save(OUT / name)


def main():
    builder = load_builder()
    save_minimal("minimal.docx")
    save_footer(builder)
    save_table(builder)
    save_headings(builder)
    save_full_variant(builder, "full-no-tables.docx", remove_tables=True)
    save_full_variant(builder, "full-no-footer.docx", remove_footer=True)


if __name__ == "__main__":
    main()
