from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PRD.md"
OUTPUT = ROOT / "docs" / "跨平台竞技手柄_PRD.docx"

BODY_FONT = "Microsoft YaHei"
DISPLAY_FONT = "Aptos Display"
TEXT = "202A28"
MUTED = "5F6B66"
HEADER_FILL = "26302E"
ALT_FILL = "F3F5F4"
BORDER = "D9D9D9"


def set_run_font(run, font_name: str, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("NEXUS X1 / PRD v1.0    ")
    set_run_font(prefix, BODY_FONT, 8, color=MUTED)

    begin_run = paragraph.add_run()
    set_run_font(begin_run, BODY_FONT, 8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    set_run_font(instruction_run, BODY_FONT, 8, color=MUTED)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    instruction_run._r.append(instruction)

    end_run = paragraph.add_run()
    set_run_font(end_run, BODY_FONT, 8, color=MUTED)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_inline_runs(paragraph, text: str, size: float = 10.2, color: str = TEXT):
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True, color=color)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Cascadia Mono", size - 0.4, color="33403C")
        else:
            run = paragraph.add_run(token)
            set_run_font(run, BODY_FONT, size, color=color)


def add_paragraph(doc, text: str, style=None, size: float = 10.2):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.22
    add_inline_runs(paragraph, text, size=size)
    return paragraph


def add_heading(doc, text: str, level: int):
    paragraph = doc.add_heading(level=level)
    paragraph.paragraph_format.keep_with_next = True
    styles = {
        1: (18, 16, 7),
        2: (14, 12, 5),
        3: (11.5, 8, 3),
    }
    size, before, after = styles[level]
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, DISPLAY_FONT, size, bold=True, color="000000")
    return paragraph


def add_bullet(doc, text: str, ordered: bool = False):
    style = "List Number" if ordered else "List Bullet"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.18
    add_inline_runs(paragraph, text)
    return paragraph


def add_numbered_paragraph(doc, label: str, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.55)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.18
    label_run = paragraph.add_run(f"{label} ")
    set_run_font(label_run, BODY_FONT, 10.2, color=TEXT)
    add_inline_runs(paragraph, text)
    return paragraph


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            value = row_values[column_index] if column_index < len(row_values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
            elif row_index % 2 == 0:
                set_cell_shading(cell, ALT_FILL)

            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            add_inline_runs(
                paragraph,
                value.replace("<br>", "\n"),
                size=8.8,
                color="FFFFFF" if row_index == 0 else TEXT,
            )
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_markdown(markdown: str):
    lines = markdown.splitlines()
    blocks = []
    index = 0

    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip() or line.strip() == "---":
            index += 1
            continue

        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            blocks.append(("table", rows))
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            blocks.append(("heading", level, heading.group(2).strip()))
            index += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", line)
        if bullet:
            blocks.append(("bullet", bullet.group(1).strip()))
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            blocks.append(("number", f"{numbered.group(1)}.", numbered.group(2).strip()))
            index += 1
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line or next_line.startswith(("#", "|", "- ")) or re.match(r"^\d+\.\s+", next_line):
                break
            paragraph_lines.append(next_line)
            index += 1
        blocks.append(("paragraph", " ".join(paragraph_lines)))

    return blocks


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.22

    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.color.rgb = RGBColor.from_string("000000")

    title = doc.styles["Title"]
    title.font.name = DISPLAY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), DISPLAY_FONT)
    title.font.size = Pt(27)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(6)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)

    for style in doc.styles:
        p_pr = getattr(style._element, "pPr", None)
        if p_pr is None:
            continue
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)


def add_cover(doc: Document, title: str):
    title_paragraph = doc.add_paragraph(style="Title")
    run = title_paragraph.add_run(title)
    set_run_font(run, DISPLAY_FONT, 27, bold=True, color="000000")

    subtitle = doc.add_paragraph(style="Subtitle")
    run = subtitle.add_run("产品需求文档 / PRD v1.0")
    set_run_font(run, BODY_FONT, 12, color=MUTED)

    summary = doc.add_paragraph()
    summary.paragraph_format.space_before = Pt(10)
    summary.paragraph_format.space_after = Pt(12)
    add_inline_runs(
        summary,
        "产品通过平台身份、配置同步与可验证输入链路，解决玩家在 PC、主机与移动端之间切换时的手感断裂。第一版明确区分无线与有线兼容边界，并将研究假设、量产验证和上市风险写进同一份决策文档。",
        size=10.6,
    )

    doc.add_page_break()


def build_document():
    markdown = SOURCE.read_text(encoding="utf-8")
    blocks = parse_markdown(markdown)

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    add_page_number(section.footer.paragraphs[0])

    title = next((block[2] for block in blocks if block[0] == "heading" and block[1] == 1), "产品需求文档")
    add_cover(doc, title)

    previous_was_heading = False
    for block in blocks:
        kind = block[0]

        if kind == "heading":
            level = block[1]
            text = block[2]
            if level == 1:
                continue
            add_heading(doc, text, min(level - 1, 3))
            previous_was_heading = True
            continue

        if kind == "paragraph":
            add_paragraph(doc, block[1])
            previous_was_heading = False
            continue

        if kind == "bullet":
            add_bullet(doc, block[1], ordered=False)
            previous_was_heading = False
            continue

        if kind == "number":
            add_numbered_paragraph(doc, block[1], block[2])
            previous_was_heading = False
            continue

        if kind == "table":
            add_table(doc, block[1])
            previous_was_heading = False

    props = doc.core_properties
    props.title = title
    props.subject = "跨平台竞技手柄产品需求文档"
    props.author = "Product Management"
    props.keywords = "PRD, controller, cross-platform, gaming peripheral"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
