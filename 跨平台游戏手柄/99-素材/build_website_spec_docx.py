from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FRAME_DIR = PROJECT_ROOT / "99-素材" / "产品帧序列"
BUILD_DIR = PROJECT_ROOT / "99-素材" / "_website_spec_build"
OUT_DIR = PROJECT_ROOT / "07-官网方案"
DOCX_PATH = OUT_DIR / "Nexus 官网需求说明书.docx"
FLOW_PATH = BUILD_DIR / "site-flow.png"
NARRATIVE_PATH = BUILD_DIR / "scroll-narrative.png"
PRODUCT_VIEWS_PATH = BUILD_DIR / "product-views.png"

FONT_REGULAR = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"

BLACK = "#000000"
WHITE = "#FFFFFF"
NAVY = "#17365D"
NAVY_DARK = "#102A43"
PALE_BLUE = "#DDEBF7"
PALE_BLUE_ALT = "#F4F8FC"
PALE_GRAY = "#F2F2F2"
GREEN = "#E2F0D9"
GREEN_DARK = "#375623"
AMBER = "#FFF2CC"
AMBER_DARK = "#7F6000"
PURPLE = "#E4DAF7"
PURPLE_DARK = "#5B3A8E"
ORANGE = "#FCE4D6"
ORANGE_DARK = "#833C0C"
GRAY_BORDER = "#D9D9D9"
GRAY_TEXT = "#595959"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for char in paragraph:
            candidate = current + char
            if text_size(draw, candidate, fnt)[0] <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = char
        lines.append(current)
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fnt: ImageFont.FreeTypeFont,
    fill: str,
    gap: int = 10,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrap_text(draw, text, fnt, max(40, x2 - x1 - 36))
    heights = [text_size(draw, line, fnt)[1] for line in lines]
    total = sum(heights) + max(0, len(lines) - 1) * gap
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width, _ = text_size(draw, line, fnt)
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=fnt, fill=fill)
        y += height + gap


def draw_node(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fill: str,
    fnt: ImageFont.FreeTypeFont,
    text_fill: str = BLACK,
    outline: str = NAVY,
    radius: int = 26,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=5)
    draw_centered_text(draw, box, text, fnt, text_fill)


def arrow_head(draw: ImageDraw.ImageDraw, tip: tuple[float, float], angle: float, color: str, size: int = 24) -> None:
    x, y = tip
    p1 = (x - size * math.cos(angle - math.pi / 6), y - size * math.sin(angle - math.pi / 6))
    p2 = (x - size * math.cos(angle + math.pi / 6), y - size * math.sin(angle + math.pi / 6))
    draw.polygon([tip, p1, p2], fill=color)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = NAVY) -> None:
    draw.line([start, end], fill=color, width=6)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    arrow_head(draw, end, angle, color)


def create_site_flow(path: Path) -> None:
    canvas = Image.new("RGB", (2200, 2500), WHITE)
    draw = ImageDraw.Draw(canvas)
    f52 = font(52)
    f42 = font(42)

    centers = [260, 580, 900, 1220, 1540, 1860, 2180]
    left = 250
    right = 1950
    nodes = [
        ("进入网站", NAVY, WHITE),
        ("单页 Hero\n产品正面", PALE_BLUE, BLACK),
        ("产品滚动外观", PALE_BLUE, BLACK),
        ("材质与光线", AMBER, BLACK),
        ("左 45°\n一键切换与配置软件", PURPLE, BLACK),
        ("右 45°\n操控与耐久", PALE_BLUE, BLACK),
        ("兼容与 SKU", PALE_BLUE, BLACK),
        ("规格与 FAQ\n背景为模糊正面图", ORANGE, BLACK),
        ("购买区\n799 元 / 京东 / 天猫", GREEN, GREEN_DARK),
    ]

    y = 90
    node_h = 190
    for i, (label, fill, text_fill) in enumerate(nodes):
        draw_node(draw, (left, y, right, y + node_h), label, fill, f52, text_fill, NAVY, 30)
        if i < len(nodes) - 1:
            arrow(draw, ((left + right) // 2, y + node_h), ((left + right) // 2, y + 230))
        y += 230

    draw.text((120, 2320), "全部内容处于同一页面；顶部固定导航按键滚动到对应章节。", font=f42, fill=GRAY_TEXT)
    canvas.save(path, quality=95)


def create_scroll_narrative(path: Path) -> None:
    canvas = Image.new("RGB", (3600, 1500), WHITE)
    draw = ImageDraw.Draw(canvas)
    f58 = font(58)
    f46 = font(46)
    f38 = font(38)

    stages = [
        ("滚动 01", "产品正面", "主标题 + 预约", NAVY, WHITE),
        ("滚动 02", "连续外观帧", "材质、灯光、角度", PALE_BLUE, BLACK),
        ("滚动 03", "材质与光线", "外观章节完成", AMBER, BLACK),
        ("滚动 04", "左 45°", "一键切换与配置软件\n文字位于右侧", PURPLE, BLACK),
        ("滚动 05", "右 45°", "操控与耐久\n文字位于左侧", PALE_BLUE, BLACK),
        ("滚动 06", "正面", "兼容与 SKU\n概念选择器", PALE_BLUE, BLACK),
        ("滚动 07", "缩小的模糊正面", "规格与 FAQ\n白字 + 金色数值", ORANGE, BLACK),
        ("购买", "品牌收束", "799 元\n京东 / 天猫", GREEN, GREEN_DARK),
    ]

    x = 90
    top = 280
    width = 390
    height = 560
    gap = 42
    for i, (small, title, body, fill, text_fill) in enumerate(stages):
        draw_node(draw, (x, top, x + width, top + height), "", fill, f46, text_fill, NAVY, 34)
        draw_centered_text(draw, (x + 20, top + 30, x + width - 20, top + 120), small, f38, GRAY_TEXT)
        draw_centered_text(draw, (x + 20, top + 130, x + width - 20, top + 260), title, f58, text_fill)
        draw_centered_text(draw, (x + 20, top + 285, x + width - 20, top + height - 40), body, f38, text_fill)
        if i < len(stages) - 1:
            arrow(draw, (x + width, top + height // 2), (x + width + gap, top + height // 2))
        x += width + gap

    draw.text((100, 100), "单页长滚动叙事", font=f58, fill=BLACK)
    draw.text((100, 1110), "页面不跳转；外观段推荐约 3 屏。桌面使用 416 帧，手机约 200 帧，最终距离在原型阶段微调。", font=f46, fill=GRAY_TEXT)
    canvas.save(path, quality=95)


def create_product_views(path: Path) -> None:
    frames = [FRAME_DIR / "f_0240.webp", FRAME_DIR / "f_0270.webp", FRAME_DIR / "f_0340.webp"]
    labels = ["正面", "左 45°", "右 45°"]
    images = [Image.open(frame).convert("RGB") for frame in frames]
    thumb_w, thumb_h = 1000, 562
    gap = 24
    canvas = Image.new("RGB", (thumb_w * 3 + gap * 4, thumb_h + 130), "#0B0F19")
    draw = ImageDraw.Draw(canvas)
    f42 = font(42)
    x = gap
    for image, label in zip(images, labels):
        image = image.resize((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        canvas.paste(image, (x, 42))
        width, _ = text_size(draw, label, f42)
        draw.text((x + (thumb_w - width) / 2, thumb_h + 62), label, font=f42, fill=WHITE)
        x += thumb_w + gap
    canvas.save(path, quality=95)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_border(cell, color: str = GRAY_BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color.replace("#", ""))


def set_cell_margins(cell, top: int = 95, start: int = 110, bottom: int = 95, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, size: float = 9.0, bold: bool = False, color: str = BLACK) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def style_paragraph(paragraph, before: float = 0, after: float = 0, line: float = 1.15, keep_next: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = BLACK,
    size: float = 9.0,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    style_paragraph(p, 0, 0, 1.05)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    centers: set[int] | None = None,
    first_center: bool = False,
    font_size: float = 9.0,
    header_size: float = 9.0,
    body_sizes: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    centers = centers or set()

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, color=WHITE, size=header_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
    set_repeat_header(table.rows[0])
    set_cant_split(table.rows[0])

    for r_idx, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for c_idx, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in centers or (first_center and c_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            size = body_sizes[c_idx] if body_sizes else font_size
            cell = row.cells[c_idx]
            set_cell_text(cell, value, size=size, align=align)
            set_cell_shading(cell, WHITE if r_idx % 2 == 0 else PALE_BLUE_ALT)
            set_cell_border(cell)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc: Document, text: str, level: int = 1, *, page_break_before: bool = False) -> None:
    p = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)


def add_body(doc: Document, text: str, *, keep_next: bool = False) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, 0, 6, 1.25, keep_next)
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, 2, 5, 1.0, True)
    run = p.add_run(text)
    set_run_font(run, size=9.5, bold=True, color=NAVY)


def add_image(doc: Document, path: Path, width: float, alt_text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, 0, 8, 1.0)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)


def add_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, 0, 0, 1.0)
    run = p.add_run("Nexus 官网需求说明书  |  v1.1  |  第 ")
    set_run_font(run, size=8.5, color=GRAY_TEXT)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    run2 = p.add_run(" 页")
    set_run_font(run2, size=8.5, color=GRAY_TEXT)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.72)
    section.right_margin = Cm(1.72)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.65)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK.replace("#", ""))
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22

    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK.replace("#", ""))
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        p_pr = style._element.get_or_add_pPr()
        for border in p_pr.findall(qn("w:pBdr")):
            p_pr.remove(border)


def build_document() -> None:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    create_site_flow(FLOW_PATH)
    create_scroll_narrative(NARRATIVE_PATH)
    create_product_views(PRODUCT_VIEWS_PATH)

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Nexus 跨平台游戏手柄官网需求说明书")
    style_paragraph(title, 0, 5, 1.0, True)

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, 0, 12, 1.0, True)
    run = subtitle.add_run("版本 v1.1  |  更新日期 2026-09-24  |  读者 产品、设计、研发、内容运营")
    set_run_font(run, size=10, color=GRAY_TEXT)

    add_body(
        doc,
        "本说明书定义 Nexus 概念官网的目标用户、单页结构、内容边界、视觉系统、动效、转化路径和验收标准。网站为一条连续长滚动页面，从产品外观开始，依次进入一键切换、操控、兼容、规格和购买。顶部锚点导航用于在同一页面内定位章节，不再跳转到独立详情页。本项目为模拟项目，页面中的功能界面、交互和电商入口不连接真实业务系统。",
        keep_next=True,
    )

    add_heading(doc, "文档信息")
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["产品名称", "Nexus 跨平台游戏手柄"],
            ["文档名称", "官网需求说明书"],
            ["版本", "v1.1"],
            ["日期", "2026-09-24"],
            ["网站范围", "单页长滚动官网"],
            ["主要受众", "22 至 30 岁高频跨平台玩家"],
            ["商业目标", "以高端外观建立兴趣，以一键切换建立购买理由，引导预约和电商购买"],
            ["项目性质", "虚构模拟项目，不连接真实支付、订单、留资或电商系统"],
            ["缺失处理", "没有确认的目标数值写「未获取」，不用行业估算填充"],
        ],
        [1.25, 5.42],
        font_size=9.2,
    )

    add_heading(doc, "变更记录")
    add_table(
        doc,
        ["版本", "日期", "变更内容", "原因"],
        [
            ["v1.1", "2026-09-24", "将首页与详情页合并为一个连续长滚动页面；新增固定锚点导航，取消页面跳转和独立详情页", "参考 Insta360 单页产品展示结构，用户希望在一个页面内完成从外观到购买的全部浏览"],
            ["v1", "2026-09-23", "初版：锁定页面结构、首页与详情页规则、视觉方向、动效、移动端和转化要求", "历史版本"],
        ],
        [0.65, 1.08, 3.22, 1.72],
        centers={0, 1},
        body_sizes=[8.8, 8.8, 8.6, 8.6],
    )

    add_heading(doc, "1 项目定位")
    add_body(
        doc,
        "结论：网站以产品外观承担第一注意力，以一键切换承担第一购买理由。用户不需要离开当前页面，向下滚动即可从外观进入功能、规格和购买；顶部固定导航允许直接跳到任一章，并始终保留预约或购买入口。",
        keep_next=True,
    )
    add_table(
        doc,
        ["维度", "定义"],
        [
            ["品牌", "Nexus"],
            ["产品", "跨平台游戏手柄"],
            ["核心用户", "每天至少在平台之间切换一次的高频玩家"],
            ["第一受众", "最终消费者"],
            ["第二受众", "渠道商与经销商的次要访客"],
            ["用户看到后最应相信", "一键切换真正省去重新配对；授权平台覆盖完整；按键、连接和续航足够可靠"],
            ["产品核心功能", "一键切换平台、配对记忆、三模连接、充电底座、霍尔或 TMR 摇杆"],
            ["商业价值", "为京东和天猫导流，形成预约兴趣与电商点击"],
            ["页面范围", "单页长滚动官网；不做独立详情页、购买页、账户页和后台系统"],
        ],
        [1.55, 5.12],
        first_center=True,
        font_size=9.1,
    )

    add_heading(doc, "2 成功指标")
    add_table(
        doc,
        ["优先级", "指标", "定义", "目标值"],
        [
            ["P0", "预约按钮点击", "用户点击首屏或导航中的预约按钮", "未获取"],
            ["P0", "电商点击", "用户点击单页购买区中的京东或天猫入口", "未获取"],
            ["P1", "产品动画完成率", "用户滚动至外观展示结束", "未获取"],
            ["P1", "功能章节到达率", "用户滚动到左 45° 与右 45° 功能章节", "未获取"],
            ["P2", "平均停留时长", "用户在单页各章节的停留时间", "未获取"],
        ],
        [0.65, 1.45, 3.55, 1.02],
        centers={0, 3},
        body_sizes=[8.7, 8.6, 8.5, 8.5],
    )
    add_body(doc, "目标数值未获取。原型完成后应先定义统计事件，再补目标值和统计周期。")

    doc.add_page_break()
    add_heading(doc, "3 页面结构")
    add_caption(doc, "图 1 网站页面结构")
    add_image(doc, FLOW_PATH, 5.65, "单页官网结构图，从产品正面依次滚动到功能章节、规格与购买区")
    add_body(
        doc,
        "网站只包含一个长滚动页面。页面上半段负责外观、品牌和预约，下半段负责一键切换、配置软件、操控、耐久、兼容、SKU、规格和购买。每一章都有稳定锚点，顶部导航只滚动，不跳转页面。",
    )

    add_heading(doc, "3.1 单页章节职责", level=2)
    add_table(
        doc,
        ["章节", "核心任务", "主要模块", "明确不做"],
        [
            ["外观章节", "用外观建立高端认知，并引导预约或继续滚动", "Hero、外观滚动、材质特写、章节过渡", "不在此段集中展开功能参数和 FAQ"],
            ["功能章节", "用一键切换、操控与兼容建立购买理由", "一键切换、配置软件、操控耐久、兼容 SKU", "不拆出独立详情页或新路由"],
            ["决策章节", "补齐规格、FAQ 与购买入口", "规格 FAQ、包装清单、购买区、页脚声明", "不做账户、支付和订单管理"],
        ],
        [0.85, 2.02, 2.35, 1.45],
        first_center=True,
        body_sizes=[8.8, 8.6, 8.5, 8.5],
    )

    add_heading(doc, "3.2 用户路径", level=2)
    add_table(
        doc,
        ["步骤", "用户看到什么", "用户下一步"],
        [
            ["1", "Hero，产品正面与主标题", "预约，或向下滚动"],
            ["2", "416 帧外观动画，手机约 200 帧", "继续滚动查看外观"],
            ["3", "产品定帧与“查看一键切换”", "在同一页面滚动到功能章节"],
            ["4", "左 45°，一键切换与配置软件", "理解核心差异"],
            ["5", "右 45°，操控与耐久", "理解可靠性与手感"],
            ["6", "正面，兼容平台与 SKU 选择", "确认自己的平台版本"],
            ["7", "模糊背景上的规格与 FAQ", "阅读决策信息"],
            ["8", "购买区，799 元与电商入口", "点击京东或天猫"],
        ],
        [0.62, 3.30, 2.75],
        centers={0},
        body_sizes=[8.7, 8.6, 8.6],
    )

    doc.add_page_break()
    add_heading(doc, "4 单页外观章节")
    add_caption(doc, "图 2 单页长滚动叙事")
    add_image(doc, NARRATIVE_PATH, 6.55, "单页长滚动叙事图，从正面连续进入左四十五度和右四十五度功能章节")

    add_heading(doc, "4.1 外观章节模块", level=2)
    add_table(
        doc,
        ["编号", "模块", "优先级", "规则"],
        [
            ["H01", "Hero 首屏", "P0", "90% 画面展示产品外观，10% 放主标题、副标题和预约；产品状态为正面"],
            ["H02", "顶部锚点导航", "P0", "初始显示 NEXUS、设计、一键切换、操控、规格和预约；点击后平滑滚动到对应章节，不跳转页面"],
            ["H03", "外观滚动", "P0", "桌面端使用 416 帧，手机端约 200 帧；支持前后滚动，不能跳帧或抖动"],
            ["H04", "材质与光线特写", "P0", "展示握把纹理、蓝紫轮廓光、金色摇杆环和少量橙红环境光"],
            ["H05", "外观章节收束", "P0", "产品回到正面定帧，按钮为“查看一键切换”，点击后滚动到同页锚点"],
            ["H06", "页脚声明", "P0", "以小字展示概念模拟声明，不与购买按钮竞争"],
        ],
        [0.62, 1.42, 0.68, 3.95],
        centers={0, 2},
        body_sizes=[8.6, 8.6, 8.4, 8.4],
    )

    add_heading(doc, "4.2 单页首屏", level=2)
    add_table(
        doc,
        ["项目", "内容", "规则"],
        [
            ["主标题", "掌控，不设边界", "高对比白色，克制排版"],
            ["副标题", "跨平台游戏手柄", "字号低于主标题，放在主标题下方"],
            ["主按钮", "预约", "点击显示“预约通道即将开放”，不执行真实留资"],
            ["产品画面", "手柄正面，90% 区域", "蓝紫轮廓光为主，金色为轻微点缀"],
            ["首屏导航", "NEXUS + 锚点菜单 + 预约", "锚点菜单保持克制；手机端可收进菜单面板"],
            ["外观动画", "约 3 屏滚动距离", "推荐值；实际在原型阶段根据节奏和手机性能微调"],
        ],
        [1.05, 2.15, 3.47],
        first_center=True,
        font_size=8.9,
    )

    add_heading(doc, "5 单页功能与决策章节", page_break_before=True)
    add_heading(doc, "5.1 同页锚点", level=2)
    add_table(
        doc,
        ["顺序", "锚点", "是否固定", "包含内容"],
        [
            ["1", "#switch 一键切换", "同页", "一键切换概念动画、配置软件和宏设置界面"],
            ["2", "#control 操控与耐久", "同页", "摇杆、按键、扳机、耐久与连接稳定性"],
            ["3", "#compatibility 兼容与 SKU", "同页", "平台兼容列表、SKU 选择与推荐说明"],
            ["4", "#specs 规格与 FAQ", "同页", "完整规格、包装清单、充电底座、授权信息和 FAQ"],
            ["5", "#purchase 购买", "同页", "日常价 799 元，京东和天猫入口"],
        ],
        [0.62, 1.35, 0.82, 3.88],
        centers={0, 2},
        body_sizes=[8.7, 8.7, 8.5, 8.5],
    )

    add_heading(doc, "5.2 功能章节顺序", level=2)
    add_table(
        doc,
        ["编号", "产品朝向", "内容", "文字位置", "优先级"],
        [
            ["S01", "外观章节结束后左 45°", "一键切换平台与配置软件", "产品右侧", "P0"],
            ["S02", "继续旋转到右 45°", "摇杆、按键、扳机与耐久", "产品左侧", "P1"],
            ["S03", "回到正面", "平台兼容与 SKU 选择", "标准图文布局", "P1"],
            ["S04", "正面缩小、偏右、轻度模糊、压暗", "完整规格与 FAQ", "白色文字与金色数值", "P2"],
            ["S05", "品牌收束", "799 元、京东与天猫入口", "同页购买区", "P0"],
        ],
        [0.62, 1.65, 2.05, 1.65, 0.70],
        centers={0, 4},
        body_sizes=[8.5, 8.4, 8.4, 8.4, 8.4],
    )

    add_heading(doc, "5.3 一键切换与配置软件", level=2)
    add_table(
        doc,
        ["项目", "内容规则"],
        [
            ["演示形式", "明确标注为概念模拟，使用分屏或可点击的平台切换流程"],
            ["流程", "按下切换 → 旧平台退出 → 新平台连接 → 状态反馈成功"],
            ["配置软件", "作为一键切换章节的子内容，展示按键映射与宏设置界面"],
            ["文字密度", "标题 + 最多 2 句说明 + 3 个卖点"],
            ["视觉", "高对比；状态切换使用蓝紫光效，按钮使用金色反馈"],
        ],
        [1.25, 5.42],
        first_center=True,
        font_size=9.0,
    )

    add_heading(doc, "5.4 操控与耐久", level=2)
    add_table(
        doc,
        ["项目", "内容规则"],
        [
            ["核心内容", "霍尔或 TMR 摇杆、ABXY 按键、霍尔扳机、连接稳定性"],
            ["文字密度", "标题 + 最多 2 句说明 + 3 个卖点"],
            ["文字位置", "产品右 45° 时放在左侧留白"],
            ["证据表达", "没有真实测试素材时明确写为概念说明，不制作虚假测试数据"],
        ],
        [1.25, 5.42],
        first_center=True,
        font_size=9.0,
    )

    add_heading(doc, "5.5 兼容与 SKU", level=2)
    add_table(
        doc,
        ["项目", "规则"],
        [
            ["平台", "PC、Switch、PlayStation、Xbox、手机"],
            ["SKU", "通用三模、PS 授权、Xbox 授权、Lightning MFI"],
            ["素材", "使用同一套 PS 风格产品画面，只切换 SKU 名称、平台图标和推荐文字"],
            ["选择器", "概念选择器，不加载不同 SKU 的真实外观图"],
            ["提示", "不额外增加说明；仅依赖页脚统一概念模拟声明"],
        ],
        [1.10, 5.57],
        first_center=True,
        font_size=9.0,
    )

    add_heading(doc, "5.6 规格与 FAQ", level=2)
    add_table(
        doc,
        ["项目", "规则"],
        [
            ["背景", "正面产品缩小并偏右，轻度模糊，压暗后作为背景"],
            ["文字", "正文纯白；价格、关键数值和按钮使用金色"],
            ["内容", "完整规格、FAQ、包装清单、充电底座、授权信息"],
            ["可读性", "背景不能影响正文对比度；手机端文字覆盖在底部安全区域"],
        ],
        [1.10, 5.57],
        first_center=True,
        font_size=9.0,
    )

    doc.add_page_break()
    add_heading(doc, "6 视觉规范")
    add_caption(doc, "图 3 产品视角参考")
    add_image(doc, PRODUCT_VIEWS_PATH, 6.55, "产品图参考，包含正面、左四十五度和右四十五度视角")

    add_heading(doc, "6.1 视觉方向", level=2)
    add_table(
        doc,
        ["维度", "规则"],
        [
            ["主风格", "Rockstar 式高对比电影感，融合 Apple 式克制排版"],
            ["参考边界", "Insta360 Luna Ultra 单页结构与锚点导航；Podium 节奏；Otherlife 整场景跟随"],
            ["画面重心", "产品永远是第一视觉；文字不覆盖产品关键细节"],
            ["光线", "蓝、紫轮廓光为主，少量橙红环境光"],
            ["金色使用", "只用于摇杆环、价格、按钮和关键数据"],
            ["禁止", "静态无动效、低对比灰字、高频闪烁、廉价促销红黄配色"],
        ],
        [1.25, 5.42],
        first_center=True,
        font_size=9.0,
    )

    add_heading(doc, "6.2 色彩建议", level=2)
    add_table(
        doc,
        ["角色", "建议色值", "使用位置"],
        [
            ["深色背景", "#090B12", "页面主要背景"],
            ["深蓝层次", "#0C1A33", "区块过渡与空间层次"],
            ["蓝紫轮廓光", "#1D74FF / #7A3CFF", "产品边缘、连接状态、滚动进度"],
            ["橙红环境光", "#E55B2A", "背景氛围，轻微使用"],
            ["金色强调", "#D4AF37", "摇杆环、价格、按钮、关键数值"],
            ["主文字", "#F4F7FB", "标题与正文"],
            ["次文字", "#AAB3C0", "说明、标签和页脚"],
        ],
        [1.35, 2.15, 3.17],
        first_center=True,
        font_size=8.8,
    )
    add_body(doc, "色值为设计建议，最终应以产品渲染帧中的实际色彩为准，保证网页与产品外观一致。")

    add_heading(doc, "6.3 文字与品牌", level=2)
    add_table(
        doc,
        ["项目", "规则"],
        [
            ["语言", "中文为主，英文只做章节装饰"],
            ["英文标签", "DESIGN / SWITCH / CONTROL / COMPATIBILITY / SPECIFICATIONS"],
            ["品牌标识", "当前使用标准字 NEXUS，等待正式 Logo"],
            ["语气", "Apple 式简短克制，避免长句和夸张形容词"],
            ["标题规则", "每章最多 1 个标题、2 句说明、3 个卖点"],
        ],
        [1.30, 5.37],
        first_center=True,
        font_size=8.9,
    )

    doc.add_page_break()
    add_heading(doc, "7 动效与交互")
    add_table(
        doc,
        ["编号", "交互", "优先级", "规则"],
        [
            ["I01", "轮廓光跟随鼠标", "P0", "移动幅度轻微，不能影响文字阅读"],
            ["I02", "按钮金色反馈", "P0", "悬停出现金色描边或光带，点击有按压反馈"],
            ["I03", "平台图标依次点亮", "P0", "一键切换动画中依次展示连接成功状态"],
            ["I04", "滚动进度线", "P0", "细线提示用户当前处于单页中的哪个章节"],
            ["I05", "预约按钮反馈", "P0", "点击后显示“预约通道即将开放”"],
            ["I06", "产品呼吸与漂浮", "P0", "只有停在 45° 或正面时轻微运动"],
            ["I07", "SKU 灯光变化", "P0", "切换 SKU 时背景光和图标变化，产品画面保持同一套外观"],
        ],
        [0.62, 1.55, 0.68, 3.82],
        centers={0, 2},
        body_sizes=[8.5, 8.5, 8.4, 8.4],
    )

    add_heading(doc, "7.1 滚动规则", level=2)
    add_table(
        doc,
        ["项目", "规则"],
        [
            ["外观动画", "桌面端 416 帧；手机端约 200 帧"],
            ["方向", "向下滚动前进，向上滚动反向播放"],
            ["边界", "首帧和末帧需要锁定，避免快速滚动产生闪烁"],
            ["章节衔接", "外观章节结束后连续进入左 45°，不重新播放正面，不触发页面加载"],
            ["锚点导航", "点击导航后平滑滚动；URL hash 只定位章节，不建立独立页面"],
            ["降级", "系统开启减少动态效果时，使用静态帧替代滚动动画"],
        ],
        [1.25, 5.42],
        first_center=True,
        font_size=9.0,
    )

    add_heading(doc, "8 响应式与性能")
    add_table(
        doc,
        ["项目", "桌面端", "手机端"],
        [
            ["首屏产品", "90% 画面，产品正面", "产品全屏为背景，文字覆盖在底部安全区域"],
            ["功能章节", "产品一侧，文字位于另一侧", "同样不采用上下分栏；文字覆盖在留白侧"],
            ["帧序列", "416 帧", "约 200 帧"],
            ["导航", "透明导航，离开首屏后保留居中锚点菜单", "顶部简化为 Logo、菜单图标与预约"],
            ["规格背景", "正面偏右、轻度模糊、压暗", "正面缩小并作为背景，正文保证对比度"],
            ["性能", "预加载关键帧，滚动时平滑播放", "优先保证流畅，必要时降低采样帧数，不改变布局"],
        ],
        [1.20, 2.70, 2.77],
        first_center=True,
        body_sizes=[8.7, 8.5, 8.5],
    )

    add_heading(doc, "9 转化与状态")
    add_table(
        doc,
        ["入口", "位置", "默认文案", "点击结果"],
        [
            ["预约", "同页顶部导航与首屏", "预约", "显示“预约通道即将开放”"],
            ["查看一键切换", "外观章节末尾", "查看一键切换", "平滑滚动到 #switch，不跳转页面"],
            ["京东购买", "同页购买区", "京东购买", "按压反馈，不执行真实跳转"],
            ["天猫购买", "同页购买区", "天猫购买", "按压反馈，不执行真实跳转"],
            ["价格", "同页购买区", "日常价 799 元", "只展示，不参与交互"],
        ],
        [1.10, 1.55, 1.55, 2.47],
        first_center=True,
        body_sizes=[8.8, 8.6, 8.6, 8.6],
    )
    add_body(doc, "电商链接在模拟阶段使用占位状态。真实 JD 与天猫链接未获取，接入前不制作可跳转链接。")

    doc.add_page_break()
    add_heading(doc, "10 验收标准")
    add_table(
        doc,
        ["编号", "验收项", "通过标准", "责任"],
        [
            ["A01", "页面数量", "只有一个长滚动页面，没有独立详情页、购买页或账户页", "产品"],
            ["A02", "外观叙事", "首屏 90% 展示产品；功能从同一页面后续章节开始", "设计"],
            ["A03", "滚动动画", "桌面 416 帧、手机约 200 帧，前后滚动稳定", "研发"],
            ["A04", "功能章节朝向", "左 45°放功能文字，右 45°放操控文字，完整规格使用模糊正面背景", "设计"],
            ["A05", "文本密度", "每个 45° 章节不超过标题、2 句说明和 3 个卖点", "内容"],
            ["A06", "视觉", "高对比电影感，蓝紫为主，橙红少量，金色仅做点缀", "设计"],
            ["A07", "移动端", "不使用上下分栏，产品全屏为背景，文字位于安全区域", "设计、研发"],
            ["A08", "导航与转化", "锚点滚动不加载新页；京东和天猫保留按钮但不真实跳转", "产品、研发"],
            ["A09", "概念声明", "网页最底部展示“部分界面、交互和动态画面为概念模拟，仅用于产品展示。”", "内容"],
        ],
        [0.62, 1.42, 4.05, 0.98],
        centers={0, 3},
        body_sizes=[8.5, 8.5, 8.3, 8.3],
    )

    add_heading(doc, "11 待制作阶段确认")
    add_table(
        doc,
        ["编号", "事项", "当前结论", "处理阶段"],
        [
            ["O01", "外观动画具体滚动距离", "推荐约 3 屏，最终看原型节奏", "线框与原型"],
            ["O02", "手机端最终采样帧数", "约 200 帧，按流畅度调整", "性能测试"],
            ["O03", "品牌 Logo", "暂用标准字 NEXUS", "视觉设计"],
            ["O04", "全部页面文案", "中文为主，英文装饰，Apple 式克制", "内容设计"],
            ["O05", "真实京东与天猫链接", "未获取", "上线前"],
            ["O06", "真实预约数据存储", "模拟项目不接入", "不做"],
        ],
        [0.62, 1.62, 3.85, 0.98],
        centers={0, 3},
        body_sizes=[8.5, 8.5, 8.3, 8.4],
    )

    add_heading(doc, "12 来源与说明", page_break_before=True)
    add_table(
        doc,
        ["类别", "说明"],
        [
            ["需求输入", "2026-09-24 网站需求澄清记录，及本项目 PRD、产品帧序列 README"],
            ["产品图像", "99-素材/产品帧序列，416 张 f_0001 至 f_0416 WebP 帧"],
            ["外部参考", "Insta360 Luna Ultra 的单页结构与锚点导航，以及 Rockstar、Apple 的高对比和克制原则"],
            ["数据性质", "所有商业指标目标值未获取；本说明书不引入外部估算"],
            ["模拟声明", "页面功能、留资、电商跳转和购买状态均为概念模拟"],
        ],
        [1.28, 5.39],
        first_center=True,
        font_size=8.9,
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
    print(DOCX_PATH)
